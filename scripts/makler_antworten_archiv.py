"""Makler-Antworten-Archiv (Chris 14.10.2026): Tabelle was/wann/wo; erledigte Mails -> Ordner Immos 2027-28."""
import sys, requests, openpyxl, json
from openpyxl.styles import Font, Alignment, PatternFill
sys.path.insert(0, "/home/bolla/workspace/scripts")
from outlook_oauth2 import refresh_access_token
H = {"Authorization": "Bearer " + refresh_access_token()}
G = "https://graph.microsoft.com/v1.0"
rows = [
 ("14.10.2026 13:50","LOEGER Immobilien","Tutzing · Fr. Löger-Hohenwarter","info@loeger.de","Erstellt gern ein Suchprofil; braucht Anschrift + Kaufpreisvorstellung","14:23 Anschrift + 750–1.250 Tsd € gesendet","erledigt"),
 ("14.10.2026 14:07","FT Immobilien 24","Unterföhring · Hr. Tortorici","kontakt@ftimmobilien24.com","In Interessentenkartei aufgenommen","keine nötig","erledigt"),
 ("14.10.2026 14:33","SEDLMAYR Immobilien","Andechs · Fr. Matthern","info@sedlmayr-immo.de","Aufnahme; 2. Mail mit Link zum Online-Suchprofil","Online-Suchprofil ausgefüllt (Weilheim +50 km, Wohnung 90–110 m², 750–1.250 Tsd €)","erledigt"),
 ("14.10.2026 14:45","Kreissparkasse München Starnberg Ebersberg","ImmobilienServiceCenter · Fr. Beißengroll","jennifer.beissengroll@kskmse.de","Aufnahme in Gesuchdatei; braucht Kontaktdaten + Preisvorstellung","17:02 Kontaktdaten (ohne Tel.) + Preisrahmen gesendet","erledigt"),
 ("14.10.2026 14:59","Schneider & Prell Immobilientreuhand","Wolfratshausen · Fr. Pfeiffer","h.pfeiffer@schneider-prell.de","Bittet um Vervollständigung der Kontaktdaten per Online-Link","Formular ausgefüllt (Anschrift, Festnetz)","erledigt"),
 ("14.10.2026 15:21","LOEGER Immobilien","Tutzing","info@loeger.de","Suchprofil erstellt, melden sich bei passendem Objekt","keine nötig","erledigt"),
 ("14.10.2026 15:40","Aigner Immobilien","München Sendling · Fr. Becker","info@aigner-immobilien.de","Braucht Telefonnummer, Postanschrift, max. Kaufpreis für Exposé-Versand","Kontaktdaten + Festnetz + Preisrahmen gesendet","erledigt"),
]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Makler-Antworten"
head = ["Eingang","Makler","Ort / Kontakt","Absender","Was sie schreiben","Unsere Reaktion","Stand"]
ws.append(head)
for r in rows: ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for w, col in zip([18,34,36,34,60,48,10], "ABCDEFG"): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    if row[6].value == "offen":
        for c in row: c.fill = PatternFill("solid", fgColor="FFF2CC")
ws.freeze_panes = "A2"
for p in ("/home/bolla/workspace/data/makler_antworten_archiv.xlsx", "/mnt/d/OneDrive/Desktop/Makler_Antworten_Archiv.xlsx"):
    wb.save(p)
json.dump([dict(zip(head, r)) for r in rows], open("/home/bolla/workspace/data/makler_antworten_archiv.json","w"), ensure_ascii=False, indent=1)

# Zielordner finden
fid = None
def find(url):
    global fid
    for f in requests.get(url, headers=H).json().get("value", []):
        if f["displayName"] == "Immos 2027-28": fid = f["id"]; return
        if f.get("childFolderCount"): find(f"{G}/me/mailFolders/{f['id']}/childFolders?$top=100")
        if fid: return
find(f"{G}/me/mailFolders?$top=100")
erledigt = ("info@loeger.de","kontakt@ftimmobilien24.com","jennifer.beissengroll@kskmse.de","info@sedlmayr-immo.de","h.pfeiffer@schneider-prell.de","info@aigner-immobilien.de")
u = f"{G}/me/mailFolders/inbox/messages?$top=100&$select=id,from&$filter=receivedDateTime ge 2026-10-14T00:00:00Z"
n = 0
for x in requests.get(u, headers=H).json()["value"]:
    if x["from"]["emailAddress"]["address"].lower() in erledigt:
        r = requests.post(f"{G}/me/messages/{x['id']}/move", headers=H, json={"destinationId": fid}); n += r.status_code == 201
print("verschoben:", n)
