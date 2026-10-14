#!/usr/bin/env python3
"""Makler-Excel (Chris 14.10.2026): Blatt 1 = Antworten, Blatt 2 = Angebote mit Wertung + Kriterien-Score (Farbskala)."""
import json, openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

rows = json.load(open("/home/bolla/workspace/data/makler_antworten_neu_rows.json"))
HEAD = "305496"
wb = openpyxl.Workbook()
# ---------- Blatt 1 ----------
ws = wb.active; ws.title = "Makler-Antworten"
cols = "ABCDEFG"
for r in sorted(rows, key=int):
    row = [rows[r].get(c, "") for c in cols]
    if row[1] == "Aigner Immobilien":
        row[5] = "Kontaktdaten + Festnetz + Preisrahmen gesendet; 14.10. abends: höfliche Absage (Etagenwohnung = No-Go, Aufzug Pflicht)"
    ws.append(row)
ws.append(["14.10.2026 15:08", "Engel & Völkers Weilheim-Schongau", "Weilheim · Fr. Weber", "gabriela.weber@engelvoelkers.com",
           "Neubauprojekt „Am Hochufer“ (615.000 €); fragt nach Kaufpreisrahmen + Feedback", "Chris prüft selbst am 15.10. – nichts tun", "offen"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=HEAD)
for w, col in zip([18, 34, 30, 34, 60, 52, 14], cols): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    if row[6].value in ("offen", "manuell nötig"):
        for c in row: c.fill = PatternFill("solid", fgColor="FFF2CC")
ws.freeze_panes = "A2"

# ---------- Blatt 2: Angebote ----------
a = wb.create_sheet("Angebote & Wertung")
head = ["Datum", "Makler", "Objekt / Lage", "Typ", "Preis €", "m²", "Aufzug", "Südterrasse", "Lage im Suchraum",
        "Gesundheit/Fitness vor Ort", "Neubau/gepflegt", "Fußbodenheizung", "Tiefgarage+Ladestation", "Keller", "ruhig",
        "Preis im Rahmen", "Größe ca. 100", "Score", "Wertung", "Notiz"]
a.append(head)
# Gewichte je Kriterium-Spalte (Index in head): J=1, ?=0,5, N=0 ; K.O.: Typ=Etagenwohnung oder Aufzug=N -> Score 0
W = {"Südterrasse": 3, "Lage im Suchraum": 3, "Gesundheit/Fitness vor Ort": 2, "Neubau/gepflegt": 2, "Fußbodenheizung": 1,
     "Tiefgarage+Ladestation": 1, "Keller": 1, "ruhig": 1, "Preis im Rahmen": 2, "Größe ca. 100": 3}
ci = {h: i + 1 for i, h in enumerate(head)}
def score_formula(r):
    parts = [f'{w}*IF({L(ci[k])}{r}="J",1,IF({L(ci[k])}{r}="N",0,0.5))' for k, w in W.items()]
    return (f'=IF(OR({L(ci["Typ"])}{r}="Etagenwohnung",{L(ci["Aufzug"])}{r}="N"),0,'
            f'ROUND(100*({"+".join(parts)})/{sum(W.values())},0))')
data = [
 ["14.10.2026", "Aigner Immobilien", "Obermenzing, 81247 München (Obj.-Nr. 43421, 4 Zi.)", "Etagenwohnung", "", 91.42, "?", "?", "J", "?", "?", "?", "?", "?", "?", None, None, None, "Absage", "Etagenwohnung = No-Go; Absage 14.10. gesendet"],
 ["14.10.2026", "Engel & Völkers Weilheim", "Neubauprojekt „Am Hochufer“, Weilheim", "?", 615000, "", "?", "?", "J", "?", "J", "?", "?", "?", "?", None, None, None, "offen", "Exposé von Chris noch zu prüfen: Penthouse? Südterrasse? Aufzug?"],
]
for i, d in enumerate(data, start=2):
    a.append(d)
    a.cell(i, ci["Preis im Rahmen"], f'=IF(E{i}="","?",IF(E{i}<=1250000,"J","N"))')
    a.cell(i, ci["Größe ca. 100"], f'=IF(F{i}="","?",IF(AND(F{i}>=90,F{i}<=115),"J","N"))')
    a.cell(i, ci["Score"], score_formula(i))
LAST = 200
for i in range(len(data) + 2, LAST + 1):  # vorbereitete Leerzeilen mit Formeln
    a.cell(i, ci["Preis im Rahmen"], f'=IF(B{i}="","",IF(E{i}="","?",IF(E{i}<=1250000,"J","N")))')
    a.cell(i, ci["Größe ca. 100"], f'=IF(B{i}="","",IF(F{i}="","?",IF(AND(F{i}>=90,F{i}<=115),"J","N")))')
    f = score_formula(i); a.cell(i, ci["Score"], f'=IF(B{i}="","",{f[1:]})')
for c in a[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=HEAD)
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
a.row_dimensions[1].height = 48
widths = [12, 26, 40, 16, 12, 8] + [11] * 10 + [9, 13, 46]
for i, w in enumerate(widths, start=1): a.column_dimensions[L(i)].width = w
thin = Side(style="thin", color="BBBBBB")
for row in a.iter_rows(min_row=2, max_row=LAST):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c.column >= 7 and c.column != 20 else "left")
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for r in range(2, LAST + 1): a.cell(r, ci["Preis €"]).number_format = '#,##0 "€"'
a.freeze_panes = "C2"
# Auswahllisten
dv_t = DataValidation(type="list", formula1='"Penthouse,Dachgeschoss,Etagenwohnung,Erdgeschoss/Garten,?"', allow_blank=True)
dv_j = DataValidation(type="list", formula1='"J,N,?"', allow_blank=True)
dv_w = DataValidation(type="list", formula1='"Interessant,Vielleicht,Absage,offen"', allow_blank=True)
for dv in (dv_t, dv_j, dv_w): a.add_data_validation(dv)
dv_t.add(f"D2:D{LAST}"); dv_w.add(f"S2:S{LAST}")
for k in ("Aufzug", "Südterrasse", "Lage im Suchraum", "Gesundheit/Fitness vor Ort", "Neubau/gepflegt", "Fußbodenheizung", "Tiefgarage+Ladestation", "Keller", "ruhig"):
    dv_j.add(f"{L(ci[k])}2:{L(ci[k])}{LAST}")
# Farbskala Score: rot (0) – gelb (50) – grün (100)
sc = L(ci["Score"])
a.conditional_formatting.add(f"{sc}2:{sc}{LAST}", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                             mid_type="num", mid_value=50, mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B"))
# J/N einfärben
for rng in (f"G2:{L(ci['Größe ca. 100'])}{LAST}",):
    a.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"J"'], fill=PatternFill("solid", bgColor="C6EFCE")))
    a.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"N"'], fill=PatternFill("solid", bgColor="FFC7CE")))
wr = f"S2:S{LAST}"
a.conditional_formatting.add(wr, CellIsRule(operator="equal", formula=['"Interessant"'], fill=PatternFill("solid", bgColor="63BE7B")))
a.conditional_formatting.add(wr, CellIsRule(operator="equal", formula=['"Vielleicht"'], fill=PatternFill("solid", bgColor="FFEB84")))
a.conditional_formatting.add(wr, CellIsRule(operator="equal", formula=['"Absage"'], fill=PatternFill("solid", bgColor="F8696B")))

# ---------- Blatt 3: Legende ----------
g = wb.create_sheet("So funktioniert's")
txt = ["Blatt „Angebote & Wertung“: pro Angebot eine Zeile, Kriterien mit J / N / ? (unbekannt) ausfüllen.",
       "Score 0–100 rechnet sich selbst: J = volle Punkte, ? = halbe Punkte, N = null.",
       "K.O.-Kriterien (Score sofort 0): Typ = Etagenwohnung oder Aufzug = N.",
       "Gewichte: " + ", ".join(f"{k} ×{w}" for k, w in W.items()),
       "Preis im Rahmen: ≤ 1.250.000 € (Rahmen 750–1.250 T€). Größe ca. 100: 90–115 m² (ohne Terrasse).",
       "Farbskala Score: rot (0) – gelb (50) – grün (100). Wertung per Auswahlliste: Interessant / Vielleicht / Absage / offen."]
for t in txt: g.append([t])
g.column_dimensions["A"].width = 130
for p in ("/home/bolla/workspace/data/makler_antworten_archiv.xlsx", "/mnt/d/OneDrive/Desktop/Makler_Antworten_Archiv_neu.xlsx"):
    wb.save(p); print("gespeichert", p)
