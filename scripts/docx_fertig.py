"""Hilfsfunktion: 'Fertig'-Button (Hyperlink auf MP /api/fertig) ans Ende eines python-docx-Dokuments (Chris 16.10.2026)."""
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import urllib.parse


def add_fertig_button(doc, job_id, titel="", base="http://127.0.0.1:18790"):
    url = f"{base}/api/fertig?id={urllib.parse.quote(job_id)}&titel={urllib.parse.quote(titel)}"
    doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    part = doc.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    for tag, val in (("w:color", "FFFFFF"),):
        e = OxmlElement(tag); e.set(qn("w:val"), val); rpr.append(e)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "40"); rpr.append(sz)
    b = OxmlElement("w:b"); rpr.append(b)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "16A34A"); rpr.append(shd)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = "   ✅  FERTIG – Bolla Bescheid sagen   "; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); p._p.append(h)
    n = doc.add_paragraph("Klick auf den grünen Knopf (Strg + Klick) – Bolla bekommt die Meldung und hakt es ab.")
    n.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in n.runs: run.font.size = Pt(10); run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def _button(doc, parent_p, url, text, fill):
    rid = doc.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "FFFFFF"); rpr.append(c)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "36"); rpr.append(sz)
    rpr.append(OxmlElement("w:b"))
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill); rpr.append(shd)
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    h.append(r); parent_p._p.append(h)


def add_entscheidungs_buttons(doc, job_id, titel="", base="http://127.0.0.1:18790", anchor=None):
    """Zwei Sammel-Knöpfe 'ALLES ÜBERNEHMEN' / 'ALLES LASSEN' (Chris 16.10.2026: bei Dokumenten mit vielen Punkten immer).
    anchor=Absatz, NACH dem eingefügt wird (sonst Dokumentende). Klick -> /api/fertig?aktion=..., Job-Detail zeigt die Entscheidung."""
    q = f"id={urllib.parse.quote(job_id)}&titel={urllib.parse.quote(titel)}&aktion="
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _button(doc, p, base + "/api/fertig?" + q + "alle_uebernehmen", "   ✔  ALLES ÜBERNEHMEN   ", "1D4ED8")
    p.add_run("      ")
    _button(doc, p, base + "/api/fertig?" + q + "alle_lassen", "   ✖  ALLES LASSEN   ", "B42318")
    n = doc.add_paragraph("Schnellweg ohne Einzelprüfung (Strg + Klick). Wer einzeln entscheidet, nutzt unten FERTIG.")
    n.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in n.runs: run.font.size = Pt(10); run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    if anchor is not None:
        anchor._p.addnext(n._p); anchor._p.addnext(p._p)
