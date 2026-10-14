import re,json,collections
from docx import Document
from docx.shared import Pt,Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
D='/home/bolla/workspace/scratch/aurora_samstag_figurenpass_20261011/'
ch=open(D+'changes.md',encoding='utf-8').read()
secs=re.split(r'(?m)^### ',ch)[1:]
cnt=collections.defaultdict(lambda: collections.Counter())
ex={}
for s in secs:
    h=s.split('\n',1)[0]
    m=re.match(r'Kap\. (\d+) · ([^·]+?) · ',h)
    if not m: continue
    k=int(m.group(1)); typ=m.group(2).strip()
    r='R1' if typ in('figurenpass','Einschub (blau)') else ('R2' if typ.startswith('fp2') else 'R3')
    cnt[k][r]+=1
    vm=re.search(r'VORHER:\s*(.*?)\n\s*\nNACHHER:\s*(.*)\Z',s,re.S)
    if vm: vm=type('M',(),{'group':lambda self,i,_a=vm.group(1).strip(),_b=vm.group(2).strip():(_a,_b)[i-1]})()
    if vm and (r,typ) not in ex and len(vm.group(1))<600 and len(vm.group(2))<1100 and vm.group(2)!=vm.group(1):
        g=re.search(r'Grund:\s*(.*)',s)
        ex[(r,typ)]=(k,g.group(1) if g else '',vm.group(1).strip(),vm.group(2).strip())
v=json.load(open(D+'panel_vorher.json')); n=json.load(open(D+'panel_nachher.json'))
def fig(x):
    vals=[m for f in x['figuren'].values() for m in f.values() if isinstance(m,(int,float))]
    return sum(vals)/len(vals) if vals else 0
def f1(x): return f"{x:.1f}".replace('.',',')
doc=Document()
st=doc.styles['Normal']; st.font.name='Aptos'; st.font.size=Pt(11)
st.paragraph_format.space_after=Pt(0); st.paragraph_format.line_spacing=1.15
for s in doc.sections: s.left_margin=s.right_margin=Cm(2); s.top_margin=s.bottom_margin=Cm(1.8)
doc.add_heading('AURORA II – Verbesserungsläufe R1–R4, Kapitel 1–24',1)
def para(t,bold=False):
    p=doc.add_paragraph(); r=p.add_run(t); r.bold=bold; p.paragraph_format.space_after=Pt(4); return p
para('Stand 14.10.2026 · Einblick in die Arbeit, nichts davon musst du tun.')
para('Was die Läufe gemacht haben',True)
for t in [
 'R1 Figuren-Pass: Leser sollen mit allen Figuren mitleiden und sich mitfreuen. Kap. 13–24: direkt im Text ergänzt (2–6 Sätze, Handlung tabu). Kap. 1–12: je Kapitel bis zu 3 blaue Vorschläge in MP – bleiben drin, wenn du nichts klickst.',
 'R2 Spannung/Humor: Fehler (Daten, Zahlen, Logik), Kanon-Glättungen und Humor-Wüsten in Kap. 13–24, ohne die neue Wärme wieder wegzuglätten.',
 'R3 Kontinuität: Nachkontrolle nach R1/R2 – Zahlen, Daten, Zeitformen, Bezüge der eingefügten Sätze.',
 'R4 Messung: Nur bewertet, nichts geändert. Pro Kapitel Spannung (1–10), Humor pro 1000 Wörter und der Schnitt aus Zuneigung/Mitleid/Mitfreude aller Figuren (1–10), jeweils vorher → nachher.']:
    p=doc.add_paragraph(t,style='List Bullet'); p.paragraph_format.space_after=Pt(3)
tot=collections.Counter()
for k in cnt: tot.update(cnt[k])
tv=[n[str(k)] for k in range(1,25)]; ov=[v[str(k)] for k in range(1,25)]
sp0=sum(x['spannung'] for x in ov)/24; sp1=sum(x['spannung'] for x in tv)/24
fg0=sum(fig(x) for x in ov)/24; fg1=sum(fig(x) for x in tv)/24
hu0=sum(x['humor_pro_1000'] for x in ov)/24; hu1=sum(x['humor_pro_1000'] for x in tv)/24
para('')
para('Ergebnis in Kürze',True)
para(f"Änderungen insgesamt: R1 {tot['R1']} · R2 {tot['R2']} · R3 {tot['R3']}. "
     f"Schnitt über 24 Kapitel: Figuren-Zuneigung/Mitleid/Mitfreude {f1(fg0)} → {f1(fg1)} · Spannung {f1(sp0)} → {f1(sp1)} · Humor/1000 Wörter {f1(hu0)} → {f1(hu1)}.")
para('')
para('Kapitel für Kapitel',True)
t=doc.add_table(rows=1,cols=7); t.style='Light Grid Accent 1'
for i,h in enumerate(['Kap.','R1','R2','R3','Spannung','Humor/1000','Figuren Ø']):
    t.rows[0].cells[i].text=h
for k in range(1,25):
    c=cnt[k]; a,b=v[str(k)],n[str(k)]
    r=t.add_row().cells
    vals=[str(k),str(c['R1']) if c['R1'] else '–',str(c['R2']) if k>=13 else '–',str(c['R3']) if k>=13 else '–',
      f"{a['spannung']} → {b['spannung']}",f"{f1(a['humor_pro_1000'])} → {f1(b['humor_pro_1000'])}",f"{f1(fig(a))} → {f1(fig(b))}"]
    for i,x in enumerate(vals): r[i].text=x
for row in t.rows:
    for c in row.cells:
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(0)
            for r in p.runs: r.font.size=Pt(10.5)
para('')
para('R1–R3 = Anzahl Änderungen je Kapitel (Kap. 1–12: blaue Vorschläge). R2/R3 liefen nur für Kap. 13–24.')
doc.add_page_break()
para('Vier Beispiele (vorher → nachher; fett = neu hinzugekommen)',True)
for key,label in [(('R1','figurenpass'),'R1 Figuren-Pass'),(('R1','Einschub (blau)'),'R1 blauer Vorschlag'),(('R2','fp2:fehler'),'R2 Fehler/Kanon'),(('R3','fp3:konsistenz'),'R3 Kontinuität')]:
    if key not in ex:
        continue
    k,g,a,b=ex[key]
    para(f'{label} · Kap. {k}',True)
    if g: para('Warum: '+g)
    para('Vorher: '+a)
    p=doc.add_paragraph(); p.add_run('Nachher: ')
    if b.startswith(a):
        p.add_run(a+' ')
        r=p.add_run(b[len(a):].strip()); r.bold=True
    else:
        p.add_run(b)
    p.paragraph_format.space_after=Pt(4)
    para('')
out='/mnt/d/OneDrive/Dokumente/Bolla/AURORA II/r1,2,3,4 Kapitel 1-24.docx'
doc.save(out); print(out, dict(tot), round(sp0,1),round(sp1,1),round(fg0,1),round(fg1,1),round(hu0,1),round(hu1,1))
