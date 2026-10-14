import re,json,collections
from docx import Document
from docx.shared import Pt,Cm
D='/home/bolla/workspace/scratch/aurora_samstag_figurenpass_20261011/'
OUT='/mnt/d/OneDrive/Dokumente/Bolla/AURORA II/r1,2,3,4 Kapitel 1-24.docx'
kap=json.load(open('/home/bolla/workspace/data/aurora2.json'))['kapitel']
ch=open(D+'changes.md',encoding='utf-8').read()
secs=[]
for s in re.split(r'(?m)^### ',ch)[1:]:
    m=re.match(r'Kap\. (\d+) · ([^·]+?) · ([^·]+?) · (\S+)',s)
    vm=re.search(r'VORHER:\s*(.*?)\n\s*\nNACHHER:\s*(.*)\Z',s,re.S)
    g=re.search(r'Grund:\s*(.*)',s)
    if m and vm: secs.append(dict(k=int(m[1]),typ=m[2].strip(),thema=m[3].strip(),v=vm[1].strip(),n=vm[2].strip(),g=g[1].strip() if g else ''))
def sent_around(text,frag):
    i=text.find(frag)
    if i<0: return None
    a=max(text.rfind(c,0,i) for c in '.!?\n')
    a=0 if a<0 else a+1
    m=re.search(r'[.!?»«…]["»«]?(\s|$)',text[i+len(frag):])
    b=i+len(frag)+(m.end() if m else 0)
    return text[a:b].strip(), i-a
def full_pair(e):
    t=kap[e['k']-1]['text']
    r=sent_around(t,e['n'])
    if not r: return None
    nach,off=r
    # Vorher = derselbe Satz mit altem Fragment
    pos=nach.find(e['n'])
    vor=nach[:pos]+e['v']+nach[pos+len(e['n']):]
    return vor,nach,e['v'],e['n']
def pick_r1(thema):
    c=[e for e in secs if e['typ']=='figurenpass' and e['thema']==thema and e['n'].startswith(e['v']) and len(e['v'])<260 and len(e['n'])<700]
    c.sort(key=lambda e:len(e['n']))
    return c[len(c)//3] if c else None
def pick_fix(prefix,thema):
    c=[e for e in secs if e['typ']==prefix and e['thema']==thema and len(e['n'])<300 and e['v']!=e['n']]
    for e in c:
        p=full_pair(e)
        if p and len(p[1])<420: return e,p
    return None,None
doc=Document()
st=doc.styles['Normal']; st.font.name='Aptos'; st.font.size=Pt(11)
st.paragraph_format.space_after=Pt(0); st.paragraph_format.line_spacing=1.15
for s in doc.sections: s.left_margin=s.right_margin=Cm(2); s.top_margin=s.bottom_margin=Cm(1.8)
def para(t='',bold=False,after=4):
    p=doc.add_paragraph(); r=p.add_run(t); r.bold=bold; p.paragraph_format.space_after=Pt(after); return p
def add_row_example(label,kapnr,grund,vorher,nachher,neu_start=None):
    para(f'{label} · Kap. {kapnr}',True,2)
    if grund: para('Warum: '+grund,after=2)
    para('Vorher: '+vorher,after=2)
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(10)
    p.add_run('Nachher: ')
    if neu_start is not None and nachher.startswith(neu_start):
        p.add_run(neu_start+' '); p.add_run(nachher[len(neu_start):].strip()).bold=True
    else: p.add_run(nachher)
doc.add_heading('AURORA II – Beispiele zu R1, R2, R3 (Kapitel 13–24)',1)
para('Vorher/Nachher jeweils als ganzer Satz bzw. Absatz aus dem Buchtext. Fett = neu hinzugekommen.')
para('R1 Figuren-Pass – ein Beispiel je Thema',True)
names={'zuneigung':'Zuneigung (Figur wird liebenswerter)','mitleid':'Mitleid (man leidet mit)','mitfreude':'Mitfreude (man freut sich mit)','humor':'Humor (warm, aus Zuneigung)'}
for th in ['zuneigung','mitleid','mitfreude','humor']:
    e=pick_r1(th)
    if e: add_row_example('R1 '+names[th],e['k'],e['g'],e['v'],e['n'],e['v'])
para('R2 Fehler/Kanon/Spannung/Humor – ein Beispiel je Thema',True)
for th in ['zahl','logik','wiederholung','stil']:
    e,p=pick_fix('fp2:fehler',th)
    if e: add_row_example(f'R2 {th}',e['k'],e['g'],p[0],p[1])
para('R3 Kontinuität',True)
for th in ['zeit','wiederholung','logik']:
    e,p=pick_fix('fp3:fehler',th)
    if e: add_row_example(f'R3 {th}',e['k'],e['g'],p[0],p[1])
doc.save(OUT.replace('Kapitel 1-24.docx','Beispiele.docx')); print('ok')
