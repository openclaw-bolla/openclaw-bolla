import re
from docx import Document
from docx.shared import Pt, RGBColor
src='/home/bolla/workspace/scratch/aurora_samstag_20261003/vorschlaege_fuer_chris.md'
out='/home/bolla/workspace/scratch/aurora_samstag_20261003/AURORA_II_Vorschlaege_Kap1-8.docx'
txt=open(src,encoding='utf-8').read()
blocks=re.split(r'^### ',txt,flags=re.M)[1:]
d=Document()
st=d.styles['Normal']; st.font.name='Aptos'; st.font.size=Pt(13)
st.paragraph_format.line_spacing=1.15; st.paragraph_format.space_after=Pt(0)
d.add_heading('AURORA II – Vorschläge zu Kap. 1–8',0)
d.add_paragraph('So gehst du vor: Pro Eintrag lesen, dann in der Zeile „Entscheidung" eintragen (ja / nein / anders: …) oder einen Word-Kommentar an die Stelle setzen. Danach sagst du Bolla Bescheid, und er baut deine Entscheidungen in die Kapitel ein.')
def lab(p,t,c=None):
    r=p.add_run(t); r.bold=True
    if c: r.font.color.rgb=RGBColor(*c)
for i,b in enumerate(blocks,1):
    head,_,body=b.partition('\n')
    head=re.sub(r'\s*\(.*?\)\s*$','',head).replace('normal:','')
    d.add_heading(f'{i}. {head}',2)
    for key,col in (('Grund',(90,90,90)),('STELLE',(160,40,40)),('VORSCHLAG',(30,110,50))):
        m=re.search(rf'^{key}:\s*(.*?)(?=\n\n[A-ZÄÖÜ]+:|\n\n(?:Grund)|\Z)',body,flags=re.M|re.S)
        if m:
            p=d.add_paragraph(); p.paragraph_format.space_before=Pt(6)
            lab(p,{'Grund':'Grund: ','STELLE':'Jetzt: ','VORSCHLAG':'Vorschlag: '}[key],col)
            p.add_run(m.group(1).strip())
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(6)
    lab(p,'Entscheidung: ',(0,0,0)); p.add_run('________________________')
d.save(out); print(len(blocks))
