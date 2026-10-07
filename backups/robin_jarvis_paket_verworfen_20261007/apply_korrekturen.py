#!/usr/bin/env python3
"""Wendet die Korrekturen aus korrekturen.json auf Kopien von Robins Word-Dateien an.

Aufruf:
  python3 apply_korrekturen.py --ordner "<Ordner mit den Originalen>" --ziel "<Ausgabeordner>"            # Trockenlauf
  python3 apply_korrekturen.py --ordner ... --ziel ... --anwenden                                          # Kopien korrigieren
  python3 apply_korrekturen.py ... --anwenden --ids d05,d12                                                # nur bestimmte IDs
  python3 apply_korrekturen.py ... --anwenden --auch-fragen                                                # auch Punkte mit modus=frage_robin
  python3 apply_korrekturen.py ... --anwenden --zusaetzlich d20,d23,l01,l02,l03,l04                       # EMPFOHLEN: alle auto-Aufträge + die freigegebenen IDs in EINEM Lauf

Die Originale werden NIE verändert. Geschrieben wird nur nach <ziel>/<name>_korrigiert.docx.
Ein Auftrag wird nur ausgeführt, wenn die Anzahl der Fundstellen exakt dem Wert "erwartet" entspricht.
Sonst wird er übersprungen und gemeldet.
Danach in Word: Strg+A, F9 (Inhalts-, Abbildungs- und Tabellenverzeichnis aktualisieren).
"""
import argparse, json, os, re, sys
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm

def alle_absaetze(d):
    """Alle Absätze inkl. Tabellenzellen, jeder Absatz nur einmal (verbundene Zellen)."""
    ps = list(d.paragraphs)
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                ps.extend(c.paragraphs)
    gesehen = set(); aus = []
    for p in ps:
        if id(p._p) in gesehen: continue
        gesehen.add(id(p._p)); aus.append(p)
    return aus

def ptext(p):
    return "".join(r.text for r in p.runs)

def ersetze_im_absatz(p, spans, ersatz_fn):
    """spans: (start, ende) im zusammengesetzten Run-Text; wird von hinten nach vorn ersetzt.
    Die Formatierung des ersten betroffenen Runs bleibt erhalten."""
    for (s, e) in sorted(spans, reverse=True):
        pos = 0; erster = True
        for r in p.runs:
            l = len(r.text); rs, re_ = pos, pos + l; pos += l
            if re_ <= s or rs >= e:
                continue
            x, y = max(s, rs) - rs, min(e, re_) - rs
            if erster:
                r.text = r.text[:x] + ersatz_fn(s, e) + r.text[y:]; erster = False
            else:
                r.text = r.text[:x] + r.text[y:]

d_ref = [None]

def finde(ps, op):
    """liefert Liste (absatz, [spans])"""
    treffer = []
    start = op.get("absatz_beginnt_mit")
    if op["typ"] == "formatvorlage":
        return [(None, [(0, 0)])]
    if op["typ"] == "tabellenzeile_entfernen":
        tb = [t for t in d_ref[0].tables if t.rows[0].cells[0].text.strip() == op["tabelle_beginnt_mit"]]
        z = [r for r in tb[0].rows if r.cells[0].text.strip() == op["zeile_beginnt_mit"]] if len(tb) == 1 else []
        return [(z[0], [(0, 0)])] if len(z) == 1 else []
    if op["typ"] == "tabellen_abstaende":
        tb = [t for t in d_ref[0].tables if t.rows[0].cells[0].text.strip() == op["tabelle_beginnt_mit"]]
        return [(tb[0], [(0, 0)])] if len(tb) == 1 else []
    for p in ps:
        if op.get("absatz_stil") and p.style.name not in op["absatz_stil"]:
            continue
        t = ptext(p)
        if start is not None and not t.startswith(start):
            continue
        if op["typ"] == "ersetzen":
            if op.get("ganzer_absatz"):
                sp = [(0, len(t))] if t == op["find"] else []
            else:
                sp = [(m.start(), m.end()) for m in re.finditer(re.escape(op["find"]), t)]
        elif op["typ"] == "regex":
            sp = [(m.start(), m.end(), m) for m in re.finditer(op["pattern"], t)]
        elif op["typ"] in ("blocksatz", "fuehrende_leerzeichen_entfernen"):
            sp = [(0, 0)] if t.startswith(start) else []
        else:
            sp = []
        if sp and op.get("erster_treffer"):
            sp = sp[:1]
        if sp:
            treffer.append((p, sp))
    return treffer

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ordner", required=True); ap.add_argument("--ziel", required=True)
    ap.add_argument("--json", default=os.path.join(os.path.dirname(os.path.abspath(__file__)), "korrekturen.json"))
    ap.add_argument("--anwenden", action="store_true"); ap.add_argument("--ids"); ap.add_argument("--auch-fragen", action="store_true")
    ap.add_argument("--zusaetzlich", help="alle auto-Aufträge PLUS diese freigegebenen IDs, z. B. d20,d23,l01,l02,l03,l04")
    a = ap.parse_args()
    cfg = json.load(open(a.json, encoding="utf-8"))
    os.makedirs(a.ziel, exist_ok=True)
    nur = set(a.ids.split(",")) if a.ids else None
    extra = set(a.zusaetzlich.split(",")) if a.zusaetzlich else set()
    unbekannt = (extra | (nur or set())) - {o["id"] for o in cfg["operationen"]}
    if unbekannt: sys.exit("Unbekannte ID(s): " + ", ".join(sorted(unbekannt)))
    docs = {}; bericht = []; ok = skip = fehl = 0
    # Formatierungs-Aufträge zuerst: ihre Absatz-Anker (Textanfang) sollen noch dem Originaltext entsprechen
    reihenfolge = sorted(cfg["operationen"], key=lambda o: 0 if o["typ"] in ("blocksatz", "fuehrende_leerzeichen_entfernen") else 1)
    for op in reihenfolge:
        if nur is not None and op["id"] not in nur: continue
        if op["modus"] == "frage_robin" and not a.auch_fragen and nur is None and op["id"] not in extra:
            bericht.append((op["id"], "ÜBERSPRUNGEN (Robin fragen)", op["beschreibung"])); skip += 1; continue
        name = cfg["dateien"][op["datei"]]
        if name not in docs: docs[name] = docx.Document(os.path.join(a.ordner, name))
        d = docs[name]; d_ref[0] = d; ps = alle_absaetze(d)
        tr = finde(ps, op)
        n = sum(len(sp) for _, sp in tr)
        if n != op["erwartet"]:
            bericht.append((op["id"], f"FEHLER: {n} Fundstellen, erwartet {op['erwartet']}", op["beschreibung"])); fehl += 1; continue
        if a.anwenden:
            for p, sp in tr:
                if op["typ"] == "tabellenzeile_entfernen":
                    p._tr.getparent().remove(p._tr)
                elif op["typ"] == "tabellen_abstaende":
                    from docx.shared import Pt
                    for ri, row in enumerate(p.rows):
                        kopf = ri > 0 and not row.cells[1].text.strip() and row.cells[0].text.strip()
                        for cell in row.cells:
                            for para in cell.paragraphs:
                                pf = para.paragraph_format
                                if kopf:
                                    pf.space_before = Pt(op["kopfzeilen_vor_pt"]); pf.space_after = Pt(op["kopfzeilen_nach_pt"])
                                else:
                                    pf.space_before = Pt(0); pf.space_after = Pt(op["zeilen_nach_pt"])
                elif op["typ"] == "formatvorlage":
                    st = d.styles[op["stil"]]; pf = st.paragraph_format
                    if "links_cm" in op: pf.left_indent = Cm(op["links_cm"])
                    if "erste_zeile_cm" in op: pf.first_line_indent = Cm(op["erste_zeile_cm"])
                    for tb in op.get("tabs", []):
                        pf.tab_stops.add_tab_stop(Cm(tb["pos_cm"]), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.SPACES)
                elif op["typ"] == "ersetzen":
                    ersetze_im_absatz(p, sp, lambda s, e: op["replace"])
                elif op["typ"] == "regex":
                    spans = [(x[0], x[1]) for x in sp]; ms = {(x[0], x[1]): x[2] for x in sp}
                    ersetze_im_absatz(p, spans, lambda s, e: ms[(s, e)].expand(op["replace"]))
                elif op["typ"] == "blocksatz":
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                elif op["typ"] == "fuehrende_leerzeichen_entfernen":
                    for r in p.runs:
                        if r.text:
                            r.text = r.text.lstrip(); break
        bericht.append((op["id"], f"OK ({n}×)" + ("" if a.anwenden else " [Trockenlauf]"), op["beschreibung"])); ok += 1
    if a.anwenden:
        for name, d in docs.items():
            out = os.path.join(a.ziel, name.replace(".docx", "_korrigiert.docx")); d.save(out); print("geschrieben:", out)
    for i, s, b in bericht: print(f"{i:5} {s:42} {b[:70]}")
    print(f"\nOK {ok} · übersprungen {skip} · Fehler {fehl}")
    if a.anwenden: print("Nächster Schritt in Word: Strg+A, F9 (Verzeichnisse aktualisieren).")
    sys.exit(1 if fehl else 0)

if __name__ == "__main__":
    main()
