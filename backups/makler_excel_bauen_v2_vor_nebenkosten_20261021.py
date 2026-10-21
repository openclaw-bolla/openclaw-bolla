#!/usr/bin/env python3
"""Makler-Excel v2 (15.10.2026): schlanke Übersicht (1 Zeile/Makler, Ampel) + Angebote mit Score + Kurz-Legende."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

HEAD = "305496"
thin = Side(style="thin", color="BBBBBB")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
FILL = {"Wartet auf Objekte": "DDEBF7", "Wartet auf Antwort": "DDEBF7", "Formular manuell": "FFF2CC",
        "Absage": "EDEDED", "Neu prüfen": "E2EFDA"}

wb = openpyxl.Workbook()

# ---------- Blatt 1: Übersicht ----------
ws = wb.active; ws.title = "Übersicht"
ws.append(["Makler", "Ort", "Kontakt", "Stand", "Nächster Schritt", "Letzte Aktion"])
rows = [
 ["Engel & Völkers Weilheim-Schongau", "Weilheim · Fr. Weber", "gabriela.weber@engelvoelkers.com", "Wartet auf Antwort",
  "Neue Objekte abwarten (PH/DG, Südterrasse, Aufzug ~100 m²)", "15.10. Hochufer abgelehnt (keine Terrasse), Kaufpreisrahmen bis 1,25 Mio genannt"],
 ["LOEGER Immobilien", "Tutzing · Fr. Löger-Hohenwarter", "info@loeger.de", "Wartet auf Objekte",
  "—", "14.10. Suchprofil erstellt, Anschrift + 750–1.250 T€ gesendet"],
 ["SEDLMAYR Immobilien", "Andechs · Fr. Matthern", "info@sedlmayr-immo.de", "Wartet auf Objekte",
  "—", "14.10. Online-Suchprofil ausgefüllt (Weilheim +50 km, 90–110 m²)"],
 ["Kreissparkasse München Starnberg Ebersberg", "ImmobilienServiceCenter · Fr. Beißengroll", "jennifer.beissengroll@kskmse.de", "Wartet auf Objekte",
  "—", "14.10. Kontaktdaten + Preisrahmen gesendet, in Gesuchdatei"],
 ["KSK München – BaufinanzierungsCenter", "München · Hr. Basile (Baufinanzierung)", "leonardo.basile@kskmse.de", "Wartet auf Objekte",
  "—", "19.10. Antwort: nimmt uns in Datenbank auf; Anschrift + Mail gesendet (ohne Telefon)"],
 ["Lieb Immobilien", "Weilheim · Hr. Lieb", "info@lieb-immobilien.de", "Wartet auf Objekte",
  "Nur noch bei Penthouse/DG mit großer Südterrasse + Aufzug melden", "19.10. Anschrift, Festnetz + Kriterien gesendet; 21.10. Exposé Peißenberg (nur überdachter Balkon) → Absage gesendet, Kontakt bleibt offen"],
 ["deltapro Immobilien", "Gilching · Fr. Götz", "info@deltapro.immobilien", "Absage",
  "—", "21.10. Antwort mit Westpark-Sendling (kein PH/keine Terrasse) + Obermenzing DG 145 m² (Preis viel zu hoch) → Absage gesendet"],
 ["Schneider & Prell Immobilientreuhand", "Wolfratshausen · Fr. Pfeiffer", "h.pfeiffer@schneider-prell.de", "Wartet auf Objekte",
  "—", "14.10. Kontaktformular ausgefüllt (Anschrift, Festnetz)"],
 ["FT Immobilien 24", "Unterföhring · Hr. Tortorici", "kontakt@ftimmobilien24.com", "Wartet auf Objekte",
  "—", "14.10. in Interessentenkartei aufgenommen"],
 ["Engel & Völkers Starnberg", "Starnberg", "contact@ (Ersatz-Adresse)", "Wartet auf Antwort",
  "—", "15.10. Voranfrage v4 an Ersatz-Adresse gesendet (alte Mail unzustellbar)"],
 ["Engel & Völkers München West", "Germering", "nur Kontaktformular", "Formular manuell",
  "Formular selbst absenden (Captcha) – Text: scratch/makler_formulare_20261014/formular_text.txt", "14.10. Formular-Versuch: Captcha"],
 ["Engel & Völkers Münchner Westen", "Gräfelfing", "nur Kontaktformular", "Formular manuell",
  "Formular selbst absenden (Captcha)", "14.10. Formular-Versuch: Captcha"],
 ["Engel & Völkers Ostallgäu – Maria Platz", "Ostallgäu", "nur Kontaktformular", "Formular manuell",
  "Formular selbst absenden (Captcha)", "14.10. Formular-Versuch: Captcha"],
 ["VON POLL IMMOBILIEN Füssen", "Füssen", "nur Kontaktformular", "Formular manuell",
  "Formular selbst absenden (Bot-Schutz 403)", "14.10. Formular-Versuch: blockt Automat"],
 ["RE/MAX Immobilien Bad Tölz", "Bad Tölz", "nur Kontaktformular", "Formular manuell",
  "Formular selbst absenden (Security Verification)", "14.10. Formular-Versuch: blockt Automat"],
 ["Aigner Immobilien", "München Sendling · Fr. Becker", "info@aigner-immobilien.de", "Absage",
  "—", "14.10. Absage gesendet (Obermenzing = Etagenwohnung, No-Go)"],
]
for r in rows: ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=12); c.fill = PatternFill("solid", fgColor=HEAD)
    c.alignment = Alignment(vertical="center")
ws.row_dimensions[1].height = 26
for w, col in zip([38, 36, 34, 20, 56, 60], "ABCDEF"): ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = BOX; c.font = Font(size=11)
    row[0].font = Font(bold=True, size=11)
    row[3].font = Font(bold=True, size=11)
    row[3].fill = PatternFill("solid", fgColor=FILL[row[3].value])
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:F{ws.max_row}"
dv = DataValidation(type="list", formula1='"Wartet auf Antwort,Wartet auf Objekte,Formular manuell,Absage,Neu prüfen"', allow_blank=True)
ws.add_data_validation(dv); dv.add(f"D2:D{ws.max_row + 20}")

# ---------- Blatt 2: Angebote ----------
a = wb.create_sheet("Angebote")
head = ["Makler", "Objekt / Lage", "Typ", "Preis €", "m²", "Score", "Wertung", "Notiz",
        "Aufzug", "Südterrasse", "Lage im Suchraum", "Gesundheit/Fitness", "Neubau/gepflegt", "Fußbodenheizung",
        "Tiefgarage+Ladestation", "Keller", "ruhig", "Preis im Rahmen", "Größe ca. 100"]
a.append(head)
ci = {h: i + 1 for i, h in enumerate(head)}
W = {"Südterrasse": 3, "Lage im Suchraum": 3, "Gesundheit/Fitness": 2, "Neubau/gepflegt": 2, "Fußbodenheizung": 1,
     "Tiefgarage+Ladestation": 1, "Keller": 1, "ruhig": 1, "Preis im Rahmen": 2, "Größe ca. 100": 3}
def score(r):
    parts = [f'{w}*IF({L(ci[k])}{r}="J",1,IF({L(ci[k])}{r}="N",0,0.5))' for k, w in W.items()]
    return (f'IF(OR({L(ci["Typ"])}{r}="Etagenwohnung",{L(ci["Aufzug"])}{r}="N",{L(ci["Südterrasse"])}{r}="N"),0,'
            f'ROUND(100*({"+".join(parts)})/{sum(W.values())},0))')
data = [
 ["Aigner Immobilien", "Obermenzing, 81247 München (Obj. 43421, 4 Zi.)", "Etagenwohnung", None, 91.42, None, "Absage",
  "Etagenwohnung = No-Go; Absage 14.10. gesendet", "?", "?", "J", "?", "?", "?", "?", "?", "?"],
 ["Engel & Völkers Weilheim", "Neubauprojekt „Am Hochufer“, Weilheim", "?", 615000, None, None, "Absage",
  "15.10. abgelehnt: keine Terrasse", "?", "N", "J", "?", "J", "?", "?", "?", "?"],
 ["Lieb Immobilien", "Neubau-DG, Wörther Str. 45, Peißenberg (5-Familienhaus, fertig Anfang 2027)", "Dachgeschoss", 599000, 122, None, "Absage",
  "21.10. Absage: nur überdachter Balkon (große Terrasse = K.O.), Aufzug unklar; sonst FBH, WP, Keller, Garage", "?", "N", "J", "?", "J", "J", "J", "J", "?"],
 ["deltapro Immobilien", "Obermenzing DG ca. 145 m², Bj. 2008/09 (+ Neubau Westpark Sendling)", "Dachgeschoss", None, 145, None, "Absage",
  "21.10. Absage: Westpark ohne PH/Terrasse, Obermenzing Preis weit über Rahmen", "?", "?", "J", "?", "N", "?", "?", "?", "?"],
]
LAST = 100
for i, d in enumerate(data, start=2):
    for j, v in enumerate(d):
        col = ci[head[j]] if j < 8 else None
    a.cell(i, ci["Makler"], d[0]); a.cell(i, ci["Objekt / Lage"], d[1]); a.cell(i, ci["Typ"], d[2])
    a.cell(i, ci["Preis €"], d[3]); a.cell(i, ci["m²"], d[4]); a.cell(i, ci["Wertung"], d[6]); a.cell(i, ci["Notiz"], d[7])
    for k, v in zip(["Aufzug", "Südterrasse", "Lage im Suchraum", "Gesundheit/Fitness", "Neubau/gepflegt", "Fußbodenheizung",
                     "Tiefgarage+Ladestation", "Keller", "ruhig"], d[8:]):
        a.cell(i, ci[k], v)
for i in range(2, LAST + 1):
    a.cell(i, ci["Preis im Rahmen"], f'=IF(A{i}="","",IF(D{i}="","?",IF(D{i}<=1250000,"J","N")))')
    a.cell(i, ci["Größe ca. 100"], f'=IF(A{i}="","",IF(E{i}="","?",IF(AND(E{i}>=90,E{i}<=115),"J","N")))')
    a.cell(i, ci["Score"], f'=IF(A{i}="","",{score(i)})')
for c in a[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=HEAD)
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
a.row_dimensions[1].height = 44
for i, w in enumerate([30, 44, 16, 13, 8, 9, 14, 44] + [12] * 11, start=1): a.column_dimensions[L(i)].width = w
for row in a.iter_rows(min_row=2, max_row=LAST):
    for c in row:
        c.border = BOX
        c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left" if c.column <= 3 or c.column == 8 else "center")
for r in range(2, LAST + 1): a.cell(r, ci["Preis €"]).number_format = '#,##0 "€"'
a.freeze_panes = "C2"
a.auto_filter.ref = f"A1:{L(len(head))}{LAST}"
# Kriterien-Spalten ab I einklappbar (gruppiert, standardmäßig sichtbar)
a.column_dimensions.group("I", L(len(head)), hidden=False, outline_level=1)
dv_t = DataValidation(type="list", formula1='"Penthouse,Dachgeschoss,Etagenwohnung,Erdgeschoss/Garten,?"', allow_blank=True)
dv_j = DataValidation(type="list", formula1='"J,N,?"', allow_blank=True)
dv_w = DataValidation(type="list", formula1='"Interessant,Vielleicht,Absage,offen"', allow_blank=True)
for d in (dv_t, dv_j, dv_w): a.add_data_validation(d)
dv_t.add(f"C2:C{LAST}"); dv_w.add(f"G2:G{LAST}")
for k in ["Aufzug", "Südterrasse", "Lage im Suchraum", "Gesundheit/Fitness", "Neubau/gepflegt", "Fußbodenheizung",
          "Tiefgarage+Ladestation", "Keller", "ruhig"]:
    dv_j.add(f"{L(ci[k])}2:{L(ci[k])}{LAST}")
a.conditional_formatting.add(f"F2:F{LAST}", ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
    mid_type="num", mid_value=50, mid_color="FFEB84", end_type="num", end_value=100, end_color="63BE7B"))
jn = f"I2:{L(len(head))}{LAST}"
a.conditional_formatting.add(jn, CellIsRule(operator="equal", formula=['"J"'], fill=PatternFill("solid", bgColor="C6EFCE")))
a.conditional_formatting.add(jn, CellIsRule(operator="equal", formula=['"N"'], fill=PatternFill("solid", bgColor="FFC7CE")))
for val, col in (("Interessant", "63BE7B"), ("Vielleicht", "FFEB84"), ("Absage", "F8696B")):
    a.conditional_formatting.add(f"G2:G{LAST}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", bgColor=col)))

# ---------- Blatt 3: Legende ----------
g = wb.create_sheet("Legende")
for t in ["Übersicht: 1 Zeile pro Makler, Spalte „Stand“ per Auswahlliste; Filter oben in jeder Spalte.",
          "Angebote: J / N / ? je Kriterium, Score 0–100 rechnet sich selbst (J = voll, ? = halb, N = 0).",
          "K.O. (Score 0): Typ = Etagenwohnung, Aufzug = N oder keine große Terrasse (Südterrasse = N; Balkon/überdachter Balkon zählt nicht).",
          "Gewichte: " + ", ".join(f"{k} ×{w}" for k, w in W.items()),
          "Preis im Rahmen: bis 1.250.000 €. Größe ca. 100: 90–115 m².",
          "Kriterien-Spalten (I–S) lassen sich über das [–] oben einklappen."]:
    g.append([t])
g.column_dimensions["A"].width = 120

for p in ("/home/bolla/workspace/data/makler_antworten_archiv.xlsx", "/mnt/d/OneDrive/Desktop/Makler_Antworten_Archiv_neu.xlsx"):
    wb.save(p); print("gespeichert", p)
